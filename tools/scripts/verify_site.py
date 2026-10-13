"""Cross-check the website data against the recalculated Excel workbook (cached values)."""
import json, sys
import openpyxl
WB = sys.argv[1]; SITE = sys.argv[2]
SHEET = {"2022": "2022", "2021": "2021", "2020": "2020", "2019b": "ספטמבר 2019",
         "2019a": "אפריל 2019", "2015": "2015", "2013": "2013", "2009": "2009"}
OUT = {0: "11", 1: "12", 2: "13", 3: "MAARIV", 4: "HAARETZ"}
def load(name):
    s = open(SITE + "/data/" + name, encoding="utf-8").read()
    return json.loads(s[s.index("=") + 1:].rstrip().rstrip(";"))
E = {e["key"]: e for e in load("elections.js")}
P = load("polls.js")
wb = openpyxl.load_workbook(WB, data_only=True)
bad = checked = 0
def same(a, b):
    if a is None or b is None: return a is None and b is None
    return abs(a - b) < 0.006
for key, sh in SHEET.items():
    ws = wb[sh]; e = E[key]; lists = {l["name"]: l for l in e["lists"]}
    rows = list(ws.iter_rows(values_only=True))
    # all-party outlet means (table starting 'בחירת ארבע')
    i = next(k for k, r in enumerate(rows) if r[0] and str(r[0]).startswith("בחירת ארבע"))
    k = i + 2
    while rows[k][0] and rows[k][0] in lists:
        name = rows[k][0]; l = lists[name]
        for j in range(5):
            d = l["polls"]["outlets"][OUT[j]]
            checked += 1
            if not same(rows[k][1 + j], d["mean"]): bad += 1; print("MEAN", key, name, OUT[j], rows[k][1 + j], d["mean"])
        checked += 1
        if not same(rows[k][6], l["polls"]["overall"]): bad += 1; print("OVERALL", key, name)
        k += 1
    # bloc table
    i = next(k for k, r in enumerate(rows) if r[0] == "גוש" and r[6] == "ממוצע סקרים")
    rev = {v: b for b, v in e["blocLabels"].items()}
    k = i + 1
    while rows[k][0] in rev:
        b = rev[rows[k][0]]
        for j in range(5):
            d = e["blocs"]["pollsByOutlet"].get(OUT[j])
            checked += 1
            if not same(rows[k][1 + j], d[b] if d else None): bad += 1; print("BLOC", key, b, OUT[j])
        for col, val in ((6, e["blocs"]["polls"][b]), (7, e["blocs"]["exit"][b]), (8, e["blocs"]["final"][b])):
            checked += 1
            if not same(rows[k][col], val): bad += 1; print("BLOC2", key, b, col, rows[k][col], val)
        k += 1
# results sheet
for r in wb["תוצאות"].iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    key = [k for k, v in SHEET.items() if v == r[0]][0]
    l = [x for x in E[key]["lists"] if x["name"] == r[1]][0]
    checked += 2
    if l["seats"] != r[2]: bad += 1; print("SEATS", key, r[1])
    if E[key]["blocLabels"][l["bloc"]] != r[3]: bad += 1; print("BLOCLBL", key, r[1])
# exit polls sheet
for r in wb["מדגמים"].iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    key = [k for k, v in SHEET.items() if v == r[0]][0]
    ex = [x for x in E[key]["exitPolls"] if x["actual"] == r[2]][0]
    checked += 1
    if ex["seats"][r[3]] != r[4]: bad += 1; print("EXIT", key, r[2], r[3])
# polls sheet: every value
ps = wb["סקרים"]
idx = {}
for p in P: idx.setdefault((p["e"], p["end"], p["publisher"]), []).append(p)
n = 0
for r in ps.iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    key = [k for k, v in SHEET.items() if v == r[0]][0]
    end = r[1].date().isoformat() if hasattr(r[1], "date") else str(r[1])
    cands = [p for p in idx.get((key, end, r[4]), []) if p["firm"] == (r[5] or "")]
    ok = any(same(p["values"].get(r[6]), r[7]) for p in cands)
    checked += 1; n += 1
    if not ok: bad += 1; print("POLL", key, end, r[4], r[6], r[7])
print("checked", checked, "mismatches", bad, "poll rows", n, "site polls x parties",
      sum(len(p["values"]) for p in P))
